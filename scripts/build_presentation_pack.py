from __future__ import annotations

from pathlib import Path
from datetime import datetime
import os, math, subprocess, textwrap, shutil

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pptx import Presentation
from pptx.util import Inches as PInches, Pt as PPt
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.dml.color import RGBColor as PRGB
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from PIL import Image, ImageOps, ImageEnhance

ROOT = Path('/root/iassetspro-ui-hotfix-20261010')
IMG = ROOT / 'docs' / 'presentation-screenshots'
OUT = ROOT / 'docs' / 'presentation-pack'
OUT.mkdir(parents=True, exist_ok=True)

DOCX = OUT / 'iAssetsPro_Full_System_Documentation.docx'
PPTX = OUT / 'iAssetsPro_Presentation_Deck.pptx'
PDF = OUT / 'iAssetsPro_Full_System_Documentation.pdf'
NOTES = OUT / 'iAssetsPro_Presentation_Speaker_Notes.md'

GREEN = '009B77'
DARK = '151A20'
NAVY = '172033'
SLATE = '5F6B7A'
LIGHT = 'F3F7F5'
GOLD = 'D79A16'
RED = 'C94747'
BLUE = '2B77B8'

images = {
    'login': IMG/'01-login.png',
    'demo': IMG/'02-demo-accounts.png',
    'dashboard': IMG/'03-dashboard.png',
    'assets': IMG/'04-assets-list.png',
    'machines': IMG/'04-machines-list.png',
    'asset_create': IMG/'annotated/05-add-asset-modal-annotated.png',
    'asset_detail': IMG/'06-asset-detail.png',
    'components': IMG/'07-asset-components-explorer.png',
    'hierarchy': IMG/'07-asset-hierarchy.png',
    'visual': IMG/'annotated/08-visual-explorer-annotated.png',
    'requests': IMG/'annotated/09-repair-requests-annotated.png',
    'workorders': IMG/'annotated/10-repair-work-orders-annotated.png',
    'materials': IMG/'annotated/11-material-requests-annotated.png',
    'tools': IMG/'annotated/12-tool-requests-annotated.png',
    'closure': IMG/'13-completion-closure.png',
    'pm_schedules': IMG/'annotated/14-pm-schedules-annotated.png',
    'pm_new': IMG/'15-pm-new-schedule.png',
    'pm_target': IMG/'annotated/16-pm-target-explorer-annotated.png',
    'pm_templates': IMG/'17-pm-templates.png',
    'pm_calendar': IMG/'18-pm-calendar.png',
}
for k,p in images.items():
    if not p.exists():
        raise SystemExit(f'Missing image {k}: {p}')

# ---------- shared helpers ----------
def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tcPr.append(shd)
    shd.set(qn('w:fill'), fill)

def set_cell_border(cell, **edges):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = tcPr.first_child_found_in('w:tcBorders')
    if tcBorders is None:
        tcBorders = OxmlElement('w:tcBorders')
        tcPr.append(tcBorders)
    for edge in ('top','left','bottom','right','insideH','insideV'):
        if edge in edges:
            tag = 'w:' + edge
            el = tcBorders.find(qn(tag))
            if el is None:
                el = OxmlElement(tag)
                tcBorders.append(el)
            for key, val in edges[edge].items():
                el.set(qn('w:' + key), str(val))

def add_doc_image(doc, path, width=6.6, caption=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Inches(width))
    if caption:
        c = doc.add_paragraph(caption)
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        c.style = doc.styles['Caption']
    return p

def add_callout(doc, title, body, color=GREEN):
    t = doc.add_table(rows=1, cols=1)
    t.autofit = True
    cell=t.cell(0,0)
    set_cell_shading(cell, 'EAF7F2' if color==GREEN else 'F2F4F7')
    set_cell_border(cell, left={'val':'single','sz':'18','color':color}, top={'val':'single','sz':'2','color':'D5DDD9'}, bottom={'val':'single','sz':'2','color':'D5DDD9'}, right={'val':'single','sz':'2','color':'D5DDD9'})
    p=cell.paragraphs[0]
    r=p.add_run(title+' '); r.bold=True; r.font.color.rgb=RGBColor.from_string(color)
    p.add_run(body)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)

def add_steps(doc, steps):
    for i, step in enumerate(steps,1):
        p=doc.add_paragraph(style='List Number')
        p.add_run(step)

def add_bullets(doc, bullets):
    for b in bullets:
        p=doc.add_paragraph(style='List Bullet')
        p.add_run(b)

def add_heading(doc, text, level=1):
    p=doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next=True
    return p

def add_status_table(doc, rows):
    table=doc.add_table(rows=1, cols=3)
    table.style='Table Grid'
    hdr=table.rows[0].cells
    for i,h in enumerate(['Area','Implemented capability','Presentation message']):
        hdr[i].text=h; set_cell_shading(hdr[i], DARK)
        for run in hdr[i].paragraphs[0].runs:
            run.font.color.rgb=RGBColor(255,255,255); run.bold=True
    for a,b,c in rows:
        cells=table.add_row().cells
        cells[0].text=a; cells[1].text=b; cells[2].text=c
        for cell in cells: cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP
    return table

# ---------- DOCX ----------
doc=Document()
sec=doc.sections[0]
sec.top_margin=Inches(0.55); sec.bottom_margin=Inches(0.55); sec.left_margin=Inches(0.62); sec.right_margin=Inches(0.62)
styles=doc.styles
styles['Normal'].font.name='Aptos'; styles['Normal'].font.size=Pt(10)
styles['Title'].font.name='Aptos Display'; styles['Title'].font.size=Pt(30); styles['Title'].font.color.rgb=RGBColor.from_string(DARK)
for s in ['Heading 1','Heading 2','Heading 3']:
    styles[s].font.name='Aptos Display'; styles[s].font.color.rgb=RGBColor.from_string(DARK)
styles['Heading 1'].font.size=Pt(21); styles['Heading 2'].font.size=Pt(16); styles['Heading 3'].font.size=Pt(12)

# cover
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.space_after=Pt(6)
r=p.add_run('iAssetsPro'); r.bold=True; r.font.size=Pt(38); r.font.color.rgb=RGBColor.from_string(GREEN)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run('Enterprise Asset Management System'); r.bold=True; r.font.size=Pt(24); r.font.color.rgb=RGBColor.from_string(DARK)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run('Full Illustrated System Documentation & Presentation Guide'); r.font.size=Pt(16); r.font.color.rgb=RGBColor.from_string(SLATE)
add_doc_image(doc, images['dashboard'], 6.8)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
p.add_run('Prepared for system presentation - October 2026').bold=True
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
p.add_run('Lightworld Technologies Ltd').font.color.rgb=RGBColor.from_string(SLATE)
doc.add_page_break()

add_heading(doc,'1. Executive Overview',1)
doc.add_paragraph('iAssetsPro is an integrated enterprise asset management platform designed to control the complete operational life of industrial assets - from registration and technical hierarchy commissioning through corrective repairs, preventive maintenance, inventory/tool custody, downtime, evidence, closure and analytics.')
add_callout(doc,'Core idea:','Every maintenance event is anchored to a physical asset and, where required, to the exact assembly, subassembly, component or part being maintained. This keeps engineering context, work history, spares, tools, costs and PM history traceable to the correct level of the machine.')
add_status_table(doc,[
('Asset register','Asset master data, plant/department/location, lifecycle and parent-asset relations','One source of truth for physical equipment.'),
('Machine hierarchy','Machine -> assembly -> subassembly -> component -> part','Maintenance can be targeted to the exact maintainable item.'),
('Repairs / RWOP','Request -> approval/conversion -> planning -> execution -> resource custody -> completion -> closure','Controlled corrective-maintenance lifecycle with auditability.'),
('Preventive maintenance','Templates, schedules, exact hierarchy targets, time/meter triggers, calendar and auto work orders','PM is planned and generated from the same asset structure used by Repairs.'),
('Materials & tools','Request, approval, issue, return, damaged/defective custody and reconciliation','Stock and tool custody remain accountable throughout work execution.'),
('Platform controls','RBAC, audit history, PostgreSQL, Redis/BullMQ background processing, automated tests','Enterprise controls protect operations and support scale.')])

add_heading(doc,'2. Access, Dashboard and Navigation',1)
add_doc_image(doc, images['login'], 6.6, 'Figure 1 - iAssetsPro login screen')
add_doc_image(doc, images['demo'], 6.6, 'Figure 2 - Demo account selector for presentation and role-based testing')
add_doc_image(doc, images['dashboard'], 6.6, 'Figure 3 - Main dashboard and module navigation')
add_bullets(doc,[
'Users authenticate into role-based workspaces. Navigation and actions are controlled by permissions and licensed modules.',
'The dashboard provides operational entry points while each specialist module exposes its own KPIs, queues and controlled actions.',
'For a presentation, begin as System Administrator or Maintenance Planner, then switch to Technician/Store roles if you want to demonstrate segregation of duties.'
])

doc.add_page_break()
add_heading(doc,'3. Asset Creation and Registration',1)
doc.add_paragraph('Asset registration creates the machine-level master record that every downstream maintenance process references. The form captures both engineering identity and organizational ownership.')
add_doc_image(doc, images['asset_create'], 6.8, 'Figure 4 - Annotated Create Asset form')
add_steps(doc,[
'Open Assets -> Machines (Asset Register) and choose Add Asset.',
'Enter identity: Name, Asset Tag, Serial Number and Category.',
'Capture engineering context: Manufacturer, Model, Year, Condition, Status and Criticality.',
'Anchor the asset physically: Location, Building, Floor, Area, Plant and Department.',
'Assign responsible personnel where appropriate and choose a Parent Asset when creating asset-to-asset hierarchy.',
'Capture description, purchase date, purchase cost and expected useful life, then Create Asset.'
])
add_callout(doc,'Why it matters:','Criticality, plant, department and status are not cosmetic fields; they become planning, filtering, prioritization and analytics dimensions across Repairs and PM.')
add_doc_image(doc, images['assets'], 6.6, 'Figure 5 - Asset register / machine list')
add_doc_image(doc, images['asset_detail'], 6.6, 'Figure 6 - Asset detail workspace')

add_heading(doc,'4. Building the Machine Technical Hierarchy',1)
doc.add_paragraph('After the machine exists, maintainable structure is commissioned beneath it. iAssetsPro distinguishes engineering hierarchy nodes so that work can be recorded at the correct physical level.')
# hierarchy table
hier=doc.add_table(rows=6, cols=3); hier.style='Table Grid'
for i,h in enumerate(['Level','Example','Maintenance meaning']):
    hier.cell(0,i).text=h; set_cell_shading(hier.cell(0,i),DARK)
    for rr in hier.cell(0,i).paragraphs[0].runs: rr.font.color.rgb=RGBColor(255,255,255); rr.bold=True
for row,(lvl,ex,meaning) in enumerate([
('Machine','Rotary Printing Machine 01','Top-level asset / parent machine'),
('Assembly','Main Drive Assembly','Major functional group'),
('Subassembly','Drive Train','Nested functional group inside an assembly'),
('Component','Main Drive Motor','Maintainable component with its own operating context'),
('Part / sub-part','Drive-End Bearing','Lowest replaceable/inspectable target')],1):
    hier.cell(row,0).text=lvl; hier.cell(row,1).text=ex; hier.cell(row,2).text=meaning
add_doc_image(doc, images['components'], 6.6, 'Figure 7 - Component/hierarchy commissioning workspace')
add_doc_image(doc, images['visual'], 6.8, 'Figure 8 - Annotated Deep Visual Explorer')
add_steps(doc,[
'Open the selected asset and use Hierarchy/Components or the Visual Explorer.',
'Create or bulk-commission assemblies and child nodes with unique component codes and parent relationships.',
'Drill from the machine through each child level; circular/unresolved parent relationships are rejected by validation.',
'Select a node to see its technical context and, where configured, engineering visuals/specifications.',
'Reuse the same hierarchy in corrective work orders, preventive schedules, part-replacement history and reports.'
])
add_callout(doc,'Presentation phrase:','“We do not only maintain a machine. We can maintain the exact bearing, motor, gearbox, shaft or subassembly inside that machine, while still preserving the relationship to the parent asset.”')

add_heading(doc,'5. Visual Engineering Explorer',1)
add_bullets(doc,[
'Left pane: hierarchy tree for deterministic drill-down.',
'Center pane: deep visual explorer with AI Realistic, Diagram, Engineering 2D and Exploded modes.',
'Right pane: selected-item context and technical specifications.',
'Zoom/reset controls allow complex equipment structures to remain navigable.',
'The explorer is reused by PM target selection so maintenance planners do not need a separate hierarchy interface.'
])

doc.add_page_break()
add_heading(doc,'6. Complete Repairs Maintenance Workflow',1)
doc.add_paragraph('The Repairs/RWOP flow controls the journey from a reported defect to a reviewed and closed work order. The same work order becomes the traceability hub for people, tasks, time, downtime, materials, tools, evidence and cost.')
add_callout(doc,'Repairs lifecycle:','Defect / Maintenance Request -> Review & approval -> Convert/Create Work Order -> Plan scope/resources -> Approve/assign -> Technician execution -> Materials/tools/downtime/evidence -> Completion report -> Supervisor review -> Close -> Analytics & history.')

add_heading(doc,'6.1 Defect or Breakdown Intake',2)
add_doc_image(doc, images['requests'], 6.8, 'Figure 9 - Annotated Maintenance Requests register')
add_steps(doc,[
'Operator, planner or authorized user raises a Maintenance Request against the affected asset.',
'The request records the issue description, priority/urgency, requester, asset and state.',
'Planner/Supervisor reviews the request and either rejects, approves or converts it to a repair work order.',
'Converted requests retain their linkage to the resulting work order for audit and history.'
])

add_heading(doc,'6.2 Work Order Planning, Approval and Assignment',2)
add_doc_image(doc, images['workorders'], 6.8, 'Figure 10 - Annotated Repair Work Orders register')
add_bullets(doc,[
'Work orders carry type, priority, status, planner, assignment and exact asset/component scope.',
'Planner can define multiple technicians and nominate a Team Leader.',
'Planned tasks, estimated duration, required materials and required tools form the execution plan.',
'Authorized approval moves the work into controlled execution rather than informal technician activity.'
])

add_heading(doc,'6.3 Technician Execution Workspace',2)
add_bullets(doc,[
'Technician starts work and the system records actual execution time through start/pause/resume/complete lifecycle actions.',
'Technicians can log task progress, shift notes, findings and work evidence/photos.',
'Assistance requests allow controlled escalation when another skill/trade is required.',
'Downtime is tracked separately so equipment unavailability can be measured and analyzed.',
'Execution history remains associated with the work order and affected hierarchy target.'
])

add_heading(doc,'6.4 Material and Spare-Part Control',2)
add_doc_image(doc, images['materials'], 6.8, 'Figure 11 - Annotated Material Requests and spare-part reconciliation')
add_steps(doc,[
'Technician or planner requests materials/spares against the work order.',
'Authorized store/planning roles review and approve requested quantities.',
'Stores issue approved quantities and inventory movement is recorded against the work order.',
'Used quantities contribute to repair cost and installed-part history.',
'Unused serviceable quantities return to usable stock.',
'Damaged or defective returned quantities do NOT become usable stock; they enter inspection/refurbishment custody for disposition.',
'Disposition can return an item to stock after QC, send it for repair/refurbishment, return it to supplier, or scrap it according to the relevant lifecycle.'
])

add_heading(doc,'6.5 Tool Request, Issue, Transfer and Return',2)
add_doc_image(doc, images['tools'], 6.8, 'Figure 12 - Annotated Tool Requests and custody state')
add_bullets(doc,[
'Tool requests are linked to the work order and quantity controlled.',
'Issue/fulfilment creates visible custody until the tool is returned/reconciled.',
'Tools can be transferred where permitted while maintaining accountability.',
'Damaged tools remain outside usable stock until repair and QC acceptance; this prevents false availability.'
])

add_heading(doc,'6.6 Completion, Review and Closure',2)
add_doc_image(doc, images['closure'], 6.6, 'Figure 13 - Completion & Closure workspace')
add_steps(doc,[
'Technician completes execution tasks and submits a completion report with findings/actions/evidence.',
'Required time logs, resource returns and completion conditions are checked before final closure.',
'Supervisor/planner reviews the completion, validates the repair outcome and closes the work order.',
'Closed work preserves the complete audit trail: request, work order, people, time, downtime, materials, tools, evidence and costs.',
'Analytics can then calculate breakdown counts, response time, repair duration, downtime, labor/material/tool cost and repeat-failure patterns.'
])

add_heading(doc,'6.7 Condition-Aware Returns and Receiving Quality Hold',2)
add_bullets(doc,[
'New purchasing receipts can be placed into quality hold when damaged/defective instead of silently increasing usable stock.',
'Quality-hold lifecycle supports quarantine -> repair/refurbish -> usable stock, supplier return or scrap.',
'Supplier-returned quantities reopen purchase-order replacement quantity where applicable.',
'Work-order material returns use the same condition-aware principle: serviceable returns credit usable stock; defective returns go to inspection/refurbishment custody.',
'Serialized spare replacement remains linked to the parent machine/component and work order for installed-part history.'
])

add_heading(doc,'7. Full Preventive Maintenance (PM) Workflow',1)
doc.add_paragraph('Preventive maintenance uses the same asset and technical hierarchy as Repairs. The objective is to define repeatable maintenance content, target the exact maintainable item, trigger work at the correct time/meter threshold and execute it through the governed work-order engine.')
add_callout(doc,'PM lifecycle:','Template/checklist -> Create schedule -> Select asset -> Drill/select exact hierarchy target -> Set time/meter frequency -> Assign priority/owner -> Activate -> Trigger/due evaluation -> Auto-generate preventive WO -> Execute checklist/resources -> Complete/close -> Update compliance and next due.')

add_heading(doc,'7.1 PM Templates and Standard Tasks',2)
add_doc_image(doc, images['pm_templates'], 6.6, 'Figure 14 - PM Templates')
add_bullets(doc,[
'Templates standardize recurring maintenance jobs and can contain multiple ordered checklist tasks.',
'Templates capture repeatable instructions, task types, expected duration and priority.',
'A schedule may reference a template so generated work orders inherit consistent execution content.'
])

add_heading(doc,'7.2 Creating a PM Schedule',2)
add_doc_image(doc, images['pm_new'], 6.6, 'Figure 15 - Create PM Schedule form')
add_steps(doc,[
'Choose New Schedule and enter the PM title/description.',
'Optionally choose a PM template to inherit standardized tasks.',
'Search all accessible assets by name or asset tag and select the parent machine.',
'Open Choose PM Target to drill into the machine hierarchy and select the exact assembly, subassembly, component or part.',
'Set frequency type and frequency value. The platform supports both time-based and meter-based schedules.',
'Set priority, estimated duration, department/assignee, lead time and Auto-generate Work Order behaviour.',
'Activate the schedule; the register then tracks next due, status and compliance context.'
])

add_heading(doc,'7.3 Exact PM Target Selection',2)
add_doc_image(doc, images['pm_target'], 6.8, 'Figure 16 - Annotated PM Target Visual Explorer')
add_callout(doc,'Important control:','Selecting a node inside the explorer is staged. The planner must choose Confirm Selection before the exact component/part is committed back to the schedule form. This avoids accidental retargeting.')

add_heading(doc,'7.4 Schedule Register, Compliance and Due State',2)
add_doc_image(doc, images['pm_schedules'], 6.8, 'Figure 17 - Annotated PM Schedules register')
add_bullets(doc,[
'PM KPIs expose schedule volume, overdue/upcoming state, generated PM work orders and completion performance.',
'Department compliance provides a planner-level health view.',
'Each schedule displays the parent asset plus its exact hierarchy target chip.',
'Time-based and meter-based PMs coexist in one register.',
'Priority, assignee and active/inactive state remain visible and auditable.',
'A View action opens fresh read-only schedule details without requiring edit permission.'
])

add_heading(doc,'7.5 PM Triggers and Automatic Work Orders',2)
add_bullets(doc,[
'Time schedules become due based on their planned cadence and lead window.',
'Meter schedules evaluate operating readings against the stored baseline and interval.',
'When a threshold is crossed, the runtime generates one preventive work order and prevents duplicate generation at the same reading.',
'The generated work order remains linked to the PM schedule, parent asset and exact component target.',
'Plant scope/RBAC prevents an unauthorized plant user from generating another plant’s PM work.',
'If a trigger-generated work order is cancelled, meter-cycle state is rearmed so the maintenance obligation is not lost.'
])

add_heading(doc,'7.6 PM Execution, Completion and Calendar',2)
add_doc_image(doc, images['pm_calendar'], 6.6, 'Figure 18 - PM Calendar')
add_bullets(doc,[
'Generated preventive work orders execute through the same controlled work-order workspace as corrective repairs.',
'Technicians perform inherited checklist tasks, record time, request materials/tools, add evidence and submit completion.',
'Closure updates schedule history and supports compliance reporting/next-due planning.',
'The PM Calendar provides a date-oriented planning view for upcoming maintenance load.'
])

add_heading(doc,'8. Roles, Governance and Traceability',1)
add_status_table(doc,[
('Operator / Requester','Raise defects and operational maintenance requests','Fast structured intake without direct maintenance-control permissions.'),
('Planner','Review requests, plan WOs, PM templates/schedules, resources','Owns work definition and maintenance planning.'),
('Supervisor','Approve/coordinate execution and review completion','Operational control and quality gate.'),
('Technician / Team Leader','Execute tasks, log time/evidence, request resources','Field execution with accountable history.'),
('Store / Inventory','Approve/issue/receive materials and tool custody','Physical inventory accountability.'),
('Administrator','Configuration, RBAC, plant/module governance','Enterprise system administration.')])
add_bullets(doc,[
'RBAC limits both visibility and mutations by role and plant scope.',
'Work records preserve linkage to the actor and lifecycle state for audit.',
'PostgreSQL is used for transactional persistence; Redis/BullMQ supports background job processing and operational queues.',
'Offline-first/sync architecture exists for supported workflows so intermittent connectivity can be handled without abandoning controlled data capture.'
])

add_heading(doc,'9. Validation and Current Delivery Status',1)
add_callout(doc,'Release quality evidence:','The PM visual-targeting/runtime release passed the complete CI chain: lint, Repairs TypeScript hard gate, full Vitest unit suite, optimized production build, production container build, platform Playwright E2E, Repairs/RWOP Playwright UAT and Required CI Gate. The live production site is responding successfully.')
add_bullets(doc,[
'The production React #130 crash in Production Resource Planning/Bottleneck Analysis was traced to malformed KPI entries and removed with regression protection.',
'Responsive dialogs now support form-aware widths and viewport containment; Purchase Order line items no longer force the modal beyond its bounds.',
'PM Schedule now searches all accessible assets, reuses the machine Visual Explorer for exact hierarchy target selection, and provides a View details flow.',
'No database schema migration was required for the PM visual-targeting upgrade because schedules already persist an exact componentId target.'
])

add_heading(doc,'10. Recommended 12-15 Minute Live Presentation Script',1)
script_steps=[
('0:00-1:00','Open with the business problem','“Industrial maintenance loses value when asset data, breakdowns, spares, tools and PM are in separate systems. iAssetsPro connects them around the physical asset.”'),
('1:00-3:00','Create/introduce an asset','Open Machines -> Add Asset. Point out identity, criticality, plant/department, parent asset and lifecycle fields.'),
('3:00-5:00','Show technical hierarchy','Open Rotary Printing Machine 01 -> Visual Explorer. Drill machine -> assembly -> subassembly -> component -> bearing part.'),
('5:00-9:00','Show corrective Repairs','Maintenance Request -> converted Work Order -> execution -> materials/tools -> completion/closure. Emphasize traceability and custody.'),
('9:00-13:00','Show PM','Open PM Schedules -> New Schedule -> choose asset -> Choose PM Target -> select exact bearing -> Confirm -> show time/meter frequency, template, auto WO and calendar.'),
('13:00-15:00','Close with control/analytics','Explain RBAC, audit trail, downtime/cost/KPI visibility, quality-hold returns and validated production release.')]
t=doc.add_table(rows=1, cols=3); t.style='Table Grid'
for i,h in enumerate(['Time','Demonstration','What to say']):
    t.cell(0,i).text=h; set_cell_shading(t.cell(0,i),DARK)
    for rr in t.cell(0,i).paragraphs[0].runs: rr.font.color.rgb=RGBColor(255,255,255); rr.bold=True
for a,b,c in script_steps:
    cells=t.add_row().cells; cells[0].text=a; cells[1].text=b; cells[2].text=c

add_heading(doc,'11. Demonstration Safety Checklist',1)
add_bullets(doc,[
'Use a prepared administrator/planner demo account and keep a technician/store account ready in another browser profile if needed.',
'Demonstrate on known sample asset “Rotary Printing Machine 01” so the hierarchy and PM examples are populated.',
'Do not delete or overwrite production sample records during presentation; create a clearly named DEMO request/schedule if live data entry is required.',
'Keep this PDF/PPTX available offline in case internet connectivity becomes unstable.',
'For a short presentation, prioritize annotated Figures 4, 8, 9, 10, 11, 12, 17 and 16.'
])

# headers/footers
for section in doc.sections:
    hp=section.header.paragraphs[0]; hp.text='iAssetsPro EAM - Illustrated System Documentation'; hp.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    hp.runs[0].font.size=Pt(8); hp.runs[0].font.color.rgb=RGBColor.from_string(SLATE)
    fp=section.footer.paragraphs[0]; fp.alignment=WD_ALIGN_PARAGRAPH.CENTER
    fp.add_run('Lightworld Technologies Ltd  |  iAssetsPro  |  October 2026').font.size=Pt(8)

doc.save(DOCX)

# ---------- PPTX ----------
prs=Presentation(); prs.slide_width=PInches(13.333); prs.slide_height=PInches(7.5)
blank=prs.slide_layouts[6]

def add_bg(slide, color='FFFFFF'):
    sh=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,0,prs.slide_width,prs.slide_height)
    sh.fill.solid(); sh.fill.fore_color.rgb=PRGB.from_string(color); sh.line.fill.background();
    slide.shapes._spTree.remove(sh._element); slide.shapes._spTree.insert(2, sh._element)

def add_title(slide, title, subtitle=None, dark=False):
    color='FFFFFF' if dark else DARK
    tb=slide.shapes.add_textbox(PInches(.55),PInches(.28),PInches(12.2),PInches(.62))
    p=tb.text_frame.paragraphs[0]; p.text=title; p.font.name='Aptos Display'; p.font.size=PPt(25); p.font.bold=True; p.font.color.rgb=PRGB.from_string(color)
    if subtitle:
        sb=slide.shapes.add_textbox(PInches(.58),PInches(.88),PInches(12),PInches(.38))
        p=sb.text_frame.paragraphs[0]; p.text=subtitle; p.font.name='Aptos'; p.font.size=PPt(11); p.font.color.rgb=PRGB.from_string('D5E3DD' if dark else SLATE)

def add_footer(slide, n):
    tx=slide.shapes.add_textbox(PInches(.55),PInches(7.17),PInches(12.2),PInches(.2))
    p=tx.text_frame.paragraphs[0]; p.text=f'iAssetsPro EAM  |  Lightworld Technologies Ltd                                           {n}'
    p.font.size=PPt(7.5); p.font.color.rgb=PRGB.from_string('7B8794')

def add_image_contain(slide, path, x,y,w,h):
    im=Image.open(path); iw,ih=im.size; ar=iw/ih; boxar=w/h
    if ar>boxar:
        ww=w; hh=w/ar; xx=x; yy=y+(h-hh)/2
    else:
        hh=h; ww=h*ar; xx=x+(w-ww)/2; yy=y
    slide.shapes.add_picture(str(path),PInches(xx),PInches(yy),PInches(ww),PInches(hh))

def add_bullet_box(slide, bullets, x=.7,y=1.4,w=5.8,h=5.4, font=17):
    box=slide.shapes.add_textbox(PInches(x),PInches(y),PInches(w),PInches(h)); tf=box.text_frame; tf.word_wrap=True; tf.margin_left=PInches(.05)
    for i,b in enumerate(bullets):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.text=b; p.level=0; p.font.name='Aptos'; p.font.size=PPt(font); p.font.color.rgb=PRGB.from_string(DARK); p.space_after=PPt(10); p.bullet=True

def slide_screenshot(title, image_key, subtitle=None):
    s=prs.slides.add_slide(blank); add_bg(s); add_title(s,title,subtitle)
    add_image_contain(s,images[image_key],.5,1.25,12.3,5.75); add_footer(s,len(prs.slides)); return s

def add_flow(slide, labels, y=3.0):
    n=len(labels); left=.45; avail=12.4; gap=.18; bw=(avail-gap*(n-1))/n
    for i,lbl in enumerate(labels):
        x=left+i*(bw+gap); sh=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,PInches(x),PInches(y),PInches(bw),PInches(1.0))
        sh.fill.solid(); sh.fill.fore_color.rgb=PRGB.from_string('EAF7F2' if i%2==0 else 'EFF3F7'); sh.line.color.rgb=PRGB.from_string(GREEN)
        tf=sh.text_frame; tf.clear(); p=tf.paragraphs[0]; p.text=lbl; p.alignment=PP_ALIGN.CENTER; p.font.size=PPt(12 if len(lbl)<20 else 10); p.font.bold=True; p.font.color.rgb=PRGB.from_string(DARK)
        if i<n-1:
            c=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,PInches(x+bw),PInches(y+.5),PInches(x+bw+gap),PInches(y+.5)); c.line.color.rgb=PRGB.from_string(GREEN); c.line.width=PPt(2)

# Slide 1
s=prs.slides.add_slide(blank); add_bg(s,DARK)
logo=s.shapes.add_textbox(PInches(.75),PInches(.85),PInches(6),PInches(.6)); p=logo.text_frame.paragraphs[0]; p.text='iAssetsPro'; p.font.size=PPt(36); p.font.bold=True; p.font.color.rgb=PRGB.from_string('19C99A')
t=s.shapes.add_textbox(PInches(.75),PInches(1.65),PInches(11.8),PInches(1.4)); p=t.text_frame.paragraphs[0]; p.text='Enterprise Asset Management System'; p.font.size=PPt(30); p.font.bold=True; p.font.color.rgb=PRGB(255,255,255)
p=t.text_frame.add_paragraph(); p.text='Asset hierarchy • Repairs maintenance • Preventive maintenance'; p.font.size=PPt(19); p.font.color.rgb=PRGB.from_string('C7D6D0')
add_image_contain(s,images['dashboard'],6.4,3.0,6.2,3.5)
add_bullet_box(s,['Illustrated system presentation','Production-validated workflows','October 2026'],.8,3.25,5.2,2.3,16)
# fix bullet text white on dark
for sh in s.shapes:
    if hasattr(sh,'text_frame') and sh.has_text_frame and sh.left < PInches(6):
        for p in sh.text_frame.paragraphs:
            if p.text and p.text not in ['iAssetsPro','Enterprise Asset Management System','Asset hierarchy • Repairs maintenance • Preventive maintenance']:
                p.font.color.rgb=PRGB.from_string('E8F1ED')
add_footer(s,1)

# 2 overview
s=prs.slides.add_slide(blank); add_bg(s); add_title(s,'What iAssetsPro Controls','One connected record from physical asset to maintenance history')
add_flow(s,['Asset master','Engineering hierarchy','Repairs / RWOP','Materials & tools','PM schedules','Analytics & audit'],2.0)
add_bullet_box(s,['Plant/department-aware asset register','Machine → assembly → subassembly → component → part','Corrective and preventive work share one governed work-order engine','Condition-aware stock/tool custody prevents false availability','RBAC, auditability and operational KPIs'],1.0,3.55,11.4,2.8,16); add_footer(s,2)

slide_screenshot('System access and role-based navigation','dashboard','Dashboard view used to enter specialist maintenance modules')
slide_screenshot('Create the parent machine / asset','asset_create','The asset master anchors engineering identity, ownership, criticality and lifecycle')

s=prs.slides.add_slide(blank); add_bg(s); add_title(s,'Commission the maintainable hierarchy','Every child remains linked to its parent machine')
add_flow(s,['MACHINE','ASSEMBLY','SUBASSEMBLY','COMPONENT','PART / SUB-PART'],2.0)
add_bullet_box(s,['Example: Rotary Printing Machine 01','Main Drive Assembly','Drive Train','Main Drive Motor','Drive-End Bearing'],1.1,3.55,11,2.3,18); add_footer(s,len(prs.slides))

slide_screenshot('Deep Visual Explorer','visual','Drill, visualize and inspect the exact engineering context of a maintainable item')

s=prs.slides.add_slide(blank); add_bg(s); add_title(s,'Corrective Repairs lifecycle','Controlled from defect intake through verified closure')
add_flow(s,['Request','Review / convert','Plan & approve','Execute','Materials / tools','Complete','Review & close'],2.15)
add_bullet_box(s,['The work order is the traceability hub','Time, downtime, people, tasks, evidence and resource usage remain attached','Closed history supports reliability and cost analytics'],1.0,3.8,11.4,2.2,17); add_footer(s,len(prs.slides))

slide_screenshot('1. Maintenance Request - controlled defect intake','requests')
slide_screenshot('2. Repair Work Order - planning and execution authority','workorders')
slide_screenshot('3. Material & spare-part custody','materials')
slide_screenshot('4. Tool custody and damaged-tool control','tools')
slide_screenshot('5. Completion & Closure','closure','Technician completion is reviewed before final work-order closure')

s=prs.slides.add_slide(blank); add_bg(s); add_title(s,'Preventive Maintenance lifecycle','Standardized work before failure')
add_flow(s,['PM template','Schedule','Select asset','Exact target','Time / meter trigger','Auto WO','Execute & close'],2.1)
add_bullet_box(s,['Same hierarchy used by Repairs','Templates standardize checklists','Schedules can be time-based or meter-based','Generated PM work orders retain parent asset + exact component target'],1.0,3.75,11.4,2.4,17); add_footer(s,len(prs.slides))

slide_screenshot('PM Templates - standardized maintenance content','pm_templates')
slide_screenshot('Create PM Schedule','pm_new','Choose cadence, priority, assignee, template and auto-generation rules')
slide_screenshot('PM exact-target Visual Explorer','pm_target','Drill to the assembly, subassembly, component or part and Confirm Selection')
slide_screenshot('PM Schedules and compliance control','pm_schedules')
slide_screenshot('PM Calendar','pm_calendar','Visual planning of upcoming preventive maintenance workload')

s=prs.slides.add_slide(blank); add_bg(s); add_title(s,'Inventory quality and maintenance custody','Condition-aware returns protect the integrity of usable stock')
add_flow(s,['Issue to WO','Use / inspect','Serviceable return','Defective quarantine','Repair / refurbish','QC accept / supplier / scrap'],2.0)
add_bullet_box(s,['Serviceable returns can re-enter usable stock','Damaged/defective material returns enter inspection custody','Damaged tools remain unavailable until repair + QC acceptance','Receiving quality hold supports supplier return and replacement quantity'],1.0,3.7,11.4,2.6,16); add_footer(s,len(prs.slides))

s=prs.slides.add_slide(blank); add_bg(s); add_title(s,'Governance and enterprise controls')
add_bullet_box(s,['Role-based access for operator, planner, supervisor, technician, store and administration','Plant-aware authorization and audit trail','PostgreSQL transactional persistence','Redis/BullMQ background processing and operational queues','Offline/sync architecture for intermittent-connectivity operations','Automated lint, unit, build, container, platform E2E and Repairs UAT release gates'],.8,1.55,11.8,4.7,18); add_footer(s,len(prs.slides))

s=prs.slides.add_slide(blank); add_bg(s,DARK); add_title(s,'Release validation', 'The demonstrated PM/runtime release passed the full Required CI Gate', dark=True)
add_bullet_box(s,['Lint ✓','Repairs TypeScript hard gate ✓','Full Vitest unit suite ✓','Optimized production build + container build ✓','Platform Playwright E2E ✓','Repairs/RWOP Playwright UAT ✓','Required CI Gate ✓'],.9,1.65,5.2,4.7,18)
for sh in s.shapes:
    if sh.has_text_frame:
        for p in sh.text_frame.paragraphs:
            if p.text and not p.text.startswith('Release validation'):
                p.font.color.rgb=PRGB.from_string('E8F1ED')
add_image_contain(s,images['pm_schedules'],6.25,1.55,6.45,4.65); add_footer(s,len(prs.slides))

s=prs.slides.add_slide(blank); add_bg(s); add_title(s,'Live demo route for tomorrow','A 12-15 minute story that shows value, not menus')
add_flow(s,['1. Add Asset','2. Drill hierarchy','3. Raise defect','4. Open WO','5. Materials/tools','6. Create PM','7. Select exact target'],2.1)
add_bullet_box(s,['Close by showing PM calendar/compliance and explaining the single traceable history from machine to part.','Keep this deck/PDF available offline as presentation backup.'],1.0,4.0,11.4,2.0,17); add_footer(s,len(prs.slides))

s=prs.slides.add_slide(blank); add_bg(s,DARK)
t=s.shapes.add_textbox(PInches(.9),PInches(2.2),PInches(11.5),PInches(1.0)); p=t.text_frame.paragraphs[0]; p.text='iAssetsPro'; p.alignment=PP_ALIGN.CENTER; p.font.size=PPt(44); p.font.bold=True; p.font.color.rgb=PRGB.from_string('19C99A')
t2=s.shapes.add_textbox(PInches(1.0),PInches(3.25),PInches(11.3),PInches(1.2)); p=t2.text_frame.paragraphs[0]; p.text='From asset structure to controlled maintenance execution.'; p.alignment=PP_ALIGN.CENTER; p.font.size=PPt(24); p.font.color.rgb=PRGB(255,255,255)
p=t2.text_frame.add_paragraph(); p.text='Thank you'; p.alignment=PP_ALIGN.CENTER; p.font.size=PPt(18); p.font.color.rgb=PRGB.from_string('C7D6D0'); add_footer(s,len(prs.slides))

prs.save(PPTX)

# ---------- speaker notes markdown ----------
NOTES.write_text('''# iAssetsPro Presentation Speaker Notes\n\n## Opening (30 seconds)\n“iAssetsPro is an enterprise asset management platform that connects the physical machine, its engineering hierarchy, corrective repairs, preventive maintenance, spares, tools, downtime, people and cost in one traceable system.”\n\n## Three ideas to repeat\n1. **Exact physical context:** Machine -> assembly -> subassembly -> component -> part.\n2. **Governed workflow:** A defect becomes an approved work order, not an informal repair.\n3. **One history:** Repairs and PM both write back to the same asset/component history.\n\n## Live demo sequence\n1. Machines -> Add Asset: show tag, serial, criticality, plant/department and parent asset.\n2. Open Rotary Printing Machine 01 -> Visual Explorer: drill to Drive-End Bearing.\n3. Repairs Maintenance -> Repair Requests: explain intake and conversion.\n4. Repair Work Orders: show planning/assignment/execution.\n5. Material Requests and Tool Requests: explain issue/return and damaged custody.\n6. Completion & Closure: explain technician report and supervisor gate.\n7. PM Maintenance -> PM Schedules -> New Schedule.\n8. Select Rotary Printing Machine 01 -> Choose PM Target -> drill to bearing -> Confirm Selection.\n9. Show time/meter frequency, PM template, auto work order and PM Calendar.\n\n## Strong closing line\n“iAssetsPro makes maintenance measurable because every request, minute, spare, tool, checklist and cost can be traced to the exact maintainable object inside the machine.”\n''', encoding='utf-8')

# ---------- convert DOCX to PDF ----------
cmd=['libreoffice','--headless','--convert-to','pdf','--outdir',str(OUT),str(DOCX)]
subprocess.run(cmd, check=True, stdout=subprocess.PIPE, stderr=subprocess.PIPE)
converted=OUT/(DOCX.stem+'.pdf')
if converted != PDF and converted.exists():
    converted.replace(PDF)

# basic validation
for f in [DOCX,PPTX,PDF,NOTES]:
    if not f.exists() or f.stat().st_size < 1000:
        raise SystemExit(f'Artifact validation failed: {f}')
print('CREATED')
for f in [DOCX,PDF,PPTX,NOTES]:
    print(f, f.stat().st_size)
